import datetime

from django import forms
from django.contrib import admin, messages
from django.utils.html import format_html
from django.utils import timezone
from ckeditor.widgets import CKEditorWidget
from import_export.admin import ExportMixin

from .resources import RegistrationResource
from .models import (
    Conference, SubTheme, Speaker, AbstractThematicArea, AbstractSubmission,
    RegistrationCategory, Registration, ProgramDay, ProgramSession, Sponsor,
    KeyMessage, ContentBlock, LOCMember, ExhibitorPackage, Exhibitor,
    ExhibitorShowcase, PaymentVerifier, AccreditationOfficer, AbstractReviewer,
    SponsorshipPackage, SponsorshipBenefit, SponsorshipPackageBenefit,
)

# Standard rich-text toolbar used across conference admin forms
_RICH = CKEditorWidget(attrs={'rows': 10}, config_name='default')
_RICH_SM = CKEditorWidget(attrs={'rows': 6}, config_name='default')


class DatePickerInput(forms.DateInput):
    input_type = 'date'


class ConferenceAdminForm(forms.ModelForm):
    class Meta:
        model = Conference
        fields = '__all__'
        widgets = {
            # Core dates
            'start_date': DatePickerInput(attrs={'placeholder': 'YYYY-MM-DD'}),
            'end_date': DatePickerInput(attrs={'placeholder': 'YYYY-MM-DD'}),
            # Key deadlines
            'abstract_deadline': DatePickerInput(attrs={'placeholder': 'YYYY-MM-DD'}),
            'early_bird_deadline': DatePickerInput(attrs={'placeholder': 'YYYY-MM-DD'}),
            'registration_deadline': DatePickerInput(attrs={'placeholder': 'YYYY-MM-DD'}),
            'notification_date': DatePickerInput(attrs={'placeholder': 'YYYY-MM-DD'}),
            'final_paper_deadline': DatePickerInput(attrs={'placeholder': 'YYYY-MM-DD'}),
            # Text fields
            'name': forms.TextInput(attrs={'placeholder': 'e.g. AKILIMO International Conference 2026'}),
            'theme': forms.TextInput(attrs={'placeholder': 'e.g. Transforming Cassava Systems for Food Security in Nigeria'}),
            'tagline': forms.TextInput(attrs={'placeholder': 'Short catchy tagline (optional)'}),
            'edition': forms.TextInput(attrs={'placeholder': 'e.g. 1st, 2nd, Inaugural'}),
            'venue': forms.TextInput(attrs={'placeholder': 'e.g. International Conference Centre'}),
            'city': forms.TextInput(attrs={'placeholder': 'e.g. Abuja'}),
            'state': forms.TextInput(attrs={'placeholder': 'e.g. FCT'}),
            'contact_email': forms.EmailInput(attrs={'placeholder': 'conference@akilimo-nigeria.org'}),
            'contact_phone': forms.TextInput(attrs={'placeholder': '+234 800 000 0000'}),
            'website_url': forms.URLInput(attrs={'placeholder': 'https://conference.akilimo-nigeria.org'}),
            # Stats bar overrides — blank means "use the live count"
            'display_registrations': forms.NumberInput(attrs={'placeholder': 'Blank = live count'}),
            'display_abstracts': forms.NumberInput(attrs={'placeholder': 'Blank = live count'}),
            'display_speakers': forms.NumberInput(attrs={'placeholder': 'Blank = live count'}),
            'display_sub_themes': forms.NumberInput(attrs={'placeholder': 'Blank = live count'}),
            'display_days': forms.NumberInput(attrs={'placeholder': 'Blank = from start/end dates'}),
            'description':       CKEditorWidget(config_name='default'),
            'objectives':        CKEditorWidget(config_name='default'),
            'expected_outcomes': CKEditorWidget(config_name='default'),
            'target_audience':   CKEditorWidget(config_name='default'),
            'key_focus_areas':   CKEditorWidget(config_name='default'),
        }


BULK_EMAIL_BATCH = 20
# Resends skip anyone emailed this recently, so a run that is interrupted
# and restarted doesn't email the same people twice.
BULK_EMAIL_RESEND_GAP = datetime.timedelta(hours=1)


def _paid(conference, mode):
    return Registration.objects.filter(
        conference=conference,
        payment_status__in=Registration.TICKET_VALID_STATUSES,
        attendance_mode=mode,
    )


def attendance_email_targets(conference, mode='new'):
    """Confirmed (or fee-waived) in-person participants not yet emailed —
    or, for a reminder, those emailed (over an hour ago) who haven't answered."""
    qs = _paid(conference, 'physical')
    if mode == 'reminder':
        return qs.filter(
            attendance_email_sent_at__lt=timezone.now() - BULK_EMAIL_RESEND_GAP,
            attendance_response='',
        )
    return qs.filter(attendance_email_sent_at__isnull=True)


def virtual_link_targets(conference, mode='new'):
    """Paid virtual attendees not yet sent the joining link — or, to resend
    (e.g. after the link changed), all of them not sent it in the last hour."""
    qs = _paid(conference, 'virtual')
    if mode == 'all':
        return qs.exclude(virtual_link_sent_at__gte=timezone.now() - BULK_EMAIL_RESEND_GAP)
    return qs.filter(virtual_link_sent_at__isnull=True)


def attendance_email_stats(conference):
    from django.db.models import Count, Q
    paid = Registration.objects.filter(
        conference=conference, payment_status__in=Registration.TICKET_VALID_STATUSES,
    )
    stats = paid.aggregate(
        eligible=Count('pk', filter=Q(attendance_mode='physical')),
        emailed=Count('pk', filter=Q(attendance_email_sent_at__isnull=False)),
        physical=Count('pk', filter=Q(attendance_response='physical')),
        virtual=Count('pk', filter=Q(attendance_response='virtual')),
        not_attending=Count('pk', filter=Q(attendance_response='not_attending')),
        no_response=Count('pk', filter=Q(attendance_email_sent_at__isnull=False, attendance_response='')),
    )
    stats['unsent'] = attendance_email_targets(conference).count()
    stats['reminder'] = attendance_email_targets(conference, 'reminder').count()
    return stats


def virtual_link_stats(conference):
    attendees = _paid(conference, 'virtual')
    return {
        'eligible': attendees.count(),
        'sent': attendees.filter(virtual_link_sent_at__isnull=False).count(),
        'unsent': virtual_link_targets(conference).count(),
        'resend': virtual_link_targets(conference, 'all').count(),
    }


# Answers to the attendance confirmation email, in the groups the responses
# page and its Excel export use. Same filters as attendance_email_stats.
ATTENDANCE_RESPONSE_GROUPS = [
    ('physical', 'Attending in person', {'attendance_response': 'physical'}),
    ('virtual', 'Attending virtually', {'attendance_response': 'virtual'}),
    ('not_attending', 'Not attending', {'attendance_response': 'not_attending'}),
    ('no_response', 'No response yet', {'attendance_email_sent_at__isnull': False, 'attendance_response': ''}),
]

ATTENDANCE_EXPORT_HEADERS = [
    'Ticket ID', 'First Name', 'Last Name', 'Email', 'Phone', 'Organization', 'Position',
    'Category', 'Response', 'Responded At', 'Email Sent At', 'Accredited',
]


def attendance_response_group(conference, filters):
    return (
        Registration.objects
        .filter(conference=conference, payment_status__in=Registration.TICKET_VALID_STATUSES, **filters)
        .select_related('category')
        .order_by('last_name', 'first_name')
    )


def _attendance_export_row(r):
    from .views import _export_cell

    def local(dt):
        return timezone.localtime(dt).replace(tzinfo=None) if dt else None
    return [_export_cell(v) for v in (
        r.ticket_id, r.first_name, r.last_name, r.email, r.phone, r.organization, r.position,
        r.category.name, r.get_attendance_response_display() or 'No response',
        local(r.attendance_responded_at), local(r.attendance_email_sent_at),
        'Yes' if r.checked_in else 'No',
    )]


def export_attendance_responses(conference, groups):
    """Excel workbook with a summary sheet and one sheet per response group."""
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    summary = wb.active
    summary.title = 'Summary'
    summary.append(['Response', 'Count'])
    for _, label, qs in groups:
        summary.append([label, qs.count()])
    summary.append(['Exported', timezone.localtime().replace(tzinfo=None)])
    summary.column_dimensions['A'].width = 24
    summary.column_dimensions['B'].width = 18
    summary['B%d' % summary.max_row].number_format = 'yyyy-mm-dd hh:mm'
    for cell in summary[1]:
        cell.font = Font(bold=True)

    for _, label, qs in groups:
        ws = wb.create_sheet(label[:31])
        ws.append(ATTENDANCE_EXPORT_HEADERS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = 'A2'
        for r in qs:
            ws.append(_attendance_export_row(r))
        for col in ('J', 'K'):
            for cell in ws[col][1:]:
                cell.number_format = 'yyyy-mm-dd hh:mm'
        for column in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in column)
            ws.column_dimensions[column[0].column_letter].width = min(max(width, 12) + 2, 50)

    stamp = timezone.localtime().strftime('%Y%m%d')
    suffix = groups[0][0] if len(groups) == 1 else 'all'
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="attendance_responses_{suffix}_{stamp}.xlsx"'
    wb.save(response)
    return response


def _send_attendance(request, registration, mode, connection):
    from .emails import send_attendance_confirmation_request
    url = request.build_absolute_uri(registration.attendance_confirmation_path())
    send_attendance_confirmation_request(registration, url, mode == 'reminder', connection)
    return 'attendance_email_sent_at'


def _send_virtual_link(request, registration, mode, connection):
    from .emails import send_virtual_joining_link
    send_virtual_joining_link(registration, connection)
    return 'virtual_link_sent_at'


# Each bulk email: who it goes to, how to send one, and which modes exist.
BULK_EMAILS = {
    'attendance': {'targets': attendance_email_targets, 'send': _send_attendance, 'modes': ('new', 'reminder')},
    'virtual-link': {'targets': virtual_link_targets, 'send': _send_virtual_link, 'modes': ('new', 'all')},
}


@admin.register(Conference)
class ConferenceAdmin(admin.ModelAdmin):
    form = ConferenceAdminForm
    list_display = ['name', 'theme', 'start_date', 'end_date', 'city', 'is_active', 'registration_open',
                    'physical_registration_open', 'virtual_registration_open', 'abstract_submission_open']
    list_filter = ['is_active', 'registration_open', 'abstract_submission_open']
    search_fields = ['name', 'theme', 'city']
    list_editable = ['is_active', 'registration_open', 'physical_registration_open',
                     'virtual_registration_open', 'abstract_submission_open']
    readonly_fields = ['stakeholder_link', 'physical_places_status', 'attendance_email_panel', 'virtual_link_panel']
    actions = ['reset_stakeholder_link']

    # ── Bulk emails (attendance confirmation, virtual joining link) ────────────

    def get_urls(self):
        from django.urls import path
        custom = [
            path('<int:pk>/bulk-email/<slug:kind>/',
                 self.admin_site.admin_view(self.bulk_email_view),
                 name='conference_conference_bulk_email'),
            path('<int:pk>/bulk-email/<slug:kind>/send/',
                 self.admin_site.admin_view(self.bulk_email_send),
                 name='conference_conference_bulk_email_send'),
            path('<int:pk>/attendance-responses/',
                 self.admin_site.admin_view(self.attendance_responses_view),
                 name='conference_conference_attendance_responses'),
        ]
        return custom + super().get_urls()

    def physical_places_status(self, obj):
        if not obj.pk:
            return "—"
        taken = obj.physical_places_taken()
        if obj.physical_capacity is None:
            return f"{taken} in-person places taken (no capacity limit set)"
        state = "FULL — in-person registration closed" if taken >= obj.physical_capacity else f"{obj.physical_capacity - taken} left"
        return f"{taken} of {obj.physical_capacity} in-person places taken · {state}"
    physical_places_status.short_description = "In-person places"

    def attendance_email_panel(self, obj):
        if not obj.pk:
            return "Save the conference first."
        from django.urls import reverse
        stats = attendance_email_stats(obj)
        return format_html(
            '<p>{} confirmed in-person participants · {} emailed · {} not yet emailed<br>'
            'Responses: {} in person · {} virtual · {} not attending · {} no response yet</p>'
            '<a class="button" href="{}" style="padding:8px 14px;">Send attendance confirmation email…</a> '
            '<a class="button" href="{}" style="padding:8px 14px;">View responses &amp; export…</a>',
            stats['eligible'], stats['emailed'], stats['unsent'],
            stats['physical'], stats['virtual'], stats['not_attending'], stats['no_response'],
            reverse('admin:conference_conference_bulk_email', args=[obj.pk, 'attendance']),
            reverse('admin:conference_conference_attendance_responses', args=[obj.pk]),
        )
    attendance_email_panel.short_description = "Attendance confirmation email"

    def virtual_link_panel(self, obj):
        if not obj.pk:
            return "Save the conference first."
        from django.urls import reverse
        stats = virtual_link_stats(obj)
        return format_html(
            '<p>{} paid virtual attendees · {} sent the link · {} not yet sent</p>'
            '<a class="button" href="{}" style="padding:8px 14px;">Send joining link…</a>',
            stats['eligible'], stats['sent'], stats['unsent'],
            reverse('admin:conference_conference_bulk_email', args=[obj.pk, 'virtual-link']),
        )
    virtual_link_panel.short_description = "Virtual joining link email"

    def _bulk_email_conference(self, request, pk, kind):
        from django.core.exceptions import PermissionDenied
        from django.http import Http404
        from django.shortcuts import get_object_or_404
        if kind not in BULK_EMAILS:
            raise Http404
        conference = get_object_or_404(Conference, pk=pk)
        if not self.has_change_permission(request, conference):
            raise PermissionDenied
        return conference

    def bulk_email_view(self, request, pk, kind):
        """Confirmation page with the send buttons and a preview of the email."""
        from django.template.response import TemplateResponse
        conference = self._bulk_email_conference(request, pk, kind)
        preview = Registration(ticket_id='ANC2026-T00000', conference=conference,
                               first_name='Ada', last_name='Obi', attendance_mode='virtual')
        if kind == 'attendance':
            stats = attendance_email_stats(conference)
            context = {
                'title': 'Send physical attendance confirmation email',
                'stats': stats,
                'buttons': [
                    ('new', f"Send to {stats['unsent']} participant(s) not yet emailed", stats['unsent']),
                    ('reminder', f"Send reminder to {stats['reminder']} who haven't responded", stats['reminder']),
                ],
                'preview_template': 'conference/emails/attendance_confirmation.html',
                'preview_subject': f"Please confirm your physical attendance — {conference.name}",
            }
        else:
            stats = virtual_link_stats(conference)
            context = {
                'title': 'Send virtual joining link',
                'stats': stats,
                'buttons': [
                    ('new', f"Send to {stats['unsent']} attendee(s) not yet sent the link", stats['unsent']),
                    ('all', f"Resend to all {stats['resend']} virtual attendees (e.g. link changed)", stats['resend']),
                ],
                'blocked': '' if conference.virtual_join_url else
                           'No joining link is set. Add the "Virtual join url" under Attendance & Venue Capacity first.',
                'preview_template': 'conference/emails/virtual_link.html',
                'preview_subject': f"Your virtual joining link — {conference.name}",
            }
        return TemplateResponse(request, 'admin/conference/conference/bulk_email.html', {
            **self.admin_site.each_context(request),
            **context,
            'kind': kind,
            'opts': self.model._meta,
            'original': conference,
            'conference': conference,
            'preview_registration': preview,
        })

    def attendance_responses_view(self, request, pk):
        """Who said they'll attend, who won't, and who hasn't answered the
        attendance confirmation email — with an Excel export."""
        from django.core.exceptions import PermissionDenied
        from django.core.paginator import Paginator
        from django.db.models import Q
        from django.shortcuts import get_object_or_404
        from django.template.response import TemplateResponse
        conference = get_object_or_404(Conference, pk=pk)
        if not self.has_view_permission(request, conference):
            raise PermissionDenied

        query = (request.GET.get('q') or '').strip()
        groups = []
        for key, label, filters in ATTENDANCE_RESPONSE_GROUPS:
            qs = attendance_response_group(conference, filters)
            for word in query.split():
                qs = qs.filter(Q(first_name__icontains=word) | Q(last_name__icontains=word)
                               | Q(email__icontains=word) | Q(organization__icontains=word)
                               | Q(ticket_id__icontains=word))
            groups.append((key, label, qs))

        status = request.GET.get('status')
        if status not in {key for key, _, _ in groups}:
            status = 'all'
        selected = groups if status == 'all' else [g for g in groups if g[0] == status]

        if request.GET.get('export') == 'xlsx':
            return export_attendance_responses(conference, selected)

        rows = [(label, r) for _, label, qs in selected for r in qs]
        page = Paginator(rows, 100).get_page(request.GET.get('page'))
        tabs = [(key, label, qs.count()) for key, label, qs in groups]
        return TemplateResponse(request, 'admin/conference/conference/attendance_responses.html', {
            **self.admin_site.each_context(request),
            'title': 'Attendance confirmation responses',
            'opts': self.model._meta,
            'original': conference,
            'conference': conference,
            'tabs': tabs,
            'total': sum(count for _, _, count in tabs),
            'status': status,
            'query': query,
            'page': page,
            'emails': ', '.join(dict.fromkeys(r.email for _, r in rows)),
            'unsent': attendance_email_targets(conference).count(),
        })

    def bulk_email_send(self, request, pk, kind):
        """Send one batch and report progress. The page calls this repeatedly
        (passing back ``after``) until ``done``, so no single request runs long
        enough to time out."""
        from django.core.mail import get_connection
        from django.http import JsonResponse, HttpResponseNotAllowed
        if request.method != 'POST':
            return HttpResponseNotAllowed(['POST'])
        conference = self._bulk_email_conference(request, pk, kind)
        config = BULK_EMAILS[kind]
        if kind == 'virtual-link' and not conference.virtual_join_url:
            return JsonResponse({'error': 'No joining link is set.'}, status=400)

        mode = request.POST.get('mode')
        if mode not in config['modes']:
            mode = 'new'
        targets = config['targets'](conference, mode)
        try:
            after = int(request.POST.get('after') or 0)
        except ValueError:
            after = 0
        batch = list(targets.filter(pk__gt=after).select_related('conference', 'category')
                     .order_by('pk')[:BULK_EMAIL_BATCH])

        sent = failed = 0
        if batch:
            with get_connection() as connection:
                for registration in batch:
                    try:
                        field = config['send'](request, registration, mode, connection)
                    except Exception:
                        failed += 1
                        continue
                    Registration.objects.filter(pk=registration.pk).update(**{field: timezone.now()})
                    sent += 1
            after = batch[-1].pk
        remaining = targets.filter(pk__gt=after).count()
        return JsonResponse({
            'sent': sent, 'failed': failed, 'after': after,
            'remaining': remaining, 'done': not batch or remaining == 0,
        })

    def stakeholder_link(self, obj):
        if not obj.pk:
            return "Save the conference first to generate the stakeholder link."
        from django.urls import reverse
        url = reverse('conference:stakeholder_register', args=[obj.stakeholder_access_token])
        return format_html(
            '<a href="{0}" target="_blank">{0}</a>'
            '<p style="color:#888;margin-top:6px;">Private, fee-free registration link. '
            'Email it directly to stakeholders — it is not linked anywhere public. '
            'Use the “Reset stakeholder registration link” action to invalidate it.</p>',
            url,
        )
    stakeholder_link.short_description = "Stakeholder registration link"

    def reset_stakeholder_link(self, request, queryset):
        import uuid
        for conference in queryset:
            conference.stakeholder_access_token = uuid.uuid4()
            conference.save(update_fields=['stakeholder_access_token'])
        self.message_user(
            request,
            f"Stakeholder registration link reset for {queryset.count()} conference(s). "
            "Any previously shared link no longer works.",
        )
    reset_stakeholder_link.short_description = "Reset stakeholder registration link"

    fieldsets = (
        ('Basic Info', {
            'fields': ('name', 'slug', 'theme', 'tagline', 'edition'),
            'description': 'Core identity of the conference.',
        }),
        ('Status & Registration', {
            'fields': ('is_active', 'registration_open', 'abstract_submission_open', 'stakeholder_link'),
            'description': 'Toggle these to open or close the conference, registration, and abstract submission.',
        }),
        ('Attendance & Venue Capacity', {
            'fields': (
                'physical_registration_open', 'virtual_registration_open',
                'physical_capacity', 'physical_places_status',
                'attendance_confirmation_deadline', 'attendance_email_panel',
                'virtual_join_url', 'virtual_join_details', 'virtual_link_panel',
            ),
            'description': (
                'Open or close in-person and virtual registration separately — the fee is the same for both. '
                'Registration is closed when neither is open (or when "Registration open" above is off).'
            ),
        }),
        ('Dates & Venue', {
            'fields': ('start_date', 'end_date', 'venue', 'city', 'state', 'country'),
        }),
        ('Key Deadlines', {
            'fields': (
                'abstract_deadline', 'early_bird_deadline',
                'notification_date', 'registration_deadline', 'final_paper_deadline',
            ),
            'description': 'Leave blank any dates that do not apply.',
        }),
        ('Content (shown on the About section of the conference page)', {
            'fields': ('description', 'target_audience', 'objectives', 'expected_outcomes', 'key_focus_areas'),
            'description': (
                'description → main paragraph | '
                'objectives → one per line | '
                'expected_outcomes → one per line | '
                'target_audience → who should attend'
            ),
        }),
        ('Contact', {
            'fields': ('contact_email', 'contact_phone', 'website_url'),
        }),
        ('Payment & Sponsorship', {
            'fields': ('bank_account_name', 'bank_account_number', 'bank_name', 'sponsorship_contact_email'),
            'description': 'Bank details shown to registrants and exhibitors who pay by bank transfer.',
        }),
        ('Media', {
            'fields': ('hero_image', 'banner_image', 'organizer_logo'),
        }),
        ('Landing Page Stats Bar', {
            'fields': (
                'display_registrations', 'display_abstracts',
                'display_speakers', 'display_sub_themes', 'display_days',
            ),
            'description': (
                'These are the five numbers in the stats bar under the hero section. '
                'Leave a field blank to show the real, live figure; type a number to '
                'show that instead.'
            ),
        }),
    )


class SubThemeAdminForm(forms.ModelForm):
    description = forms.CharField(required=False, widget=CKEditorWidget(config_name='default'))
    class Meta:
        model = SubTheme
        fields = '__all__'


@admin.register(SubTheme)
class SubThemeAdmin(admin.ModelAdmin):
    form = SubThemeAdminForm
    list_display = ['title', 'conference', 'order', 'is_active']
    list_filter = ['conference', 'is_active']
    list_editable = ['order', 'is_active']
    ordering = ['conference', 'order']


class SpeakerAdminForm(forms.ModelForm):
    bio = forms.CharField(required=False, widget=CKEditorWidget(config_name='default'))
    class Meta:
        model = Speaker
        fields = '__all__'


@admin.register(Speaker)
class SpeakerAdmin(admin.ModelAdmin):
    form = SpeakerAdminForm
    list_display = ['full_name', 'organization', 'speaker_type', 'topic', 'is_featured', 'is_active', 'photo_preview']
    list_filter = ['conference', 'speaker_type', 'is_featured', 'is_active']
    search_fields = ['full_name', 'organization', 'topic']
    list_editable = ['speaker_type', 'is_featured', 'is_active']
    ordering = ['conference', 'order']

    def photo_preview(self, obj):
        if obj.photo:
            return format_html('<img src="{}" style="width:40px;height:40px;border-radius:50%;object-fit:cover;">', obj.photo.url)
        return "—"
    photo_preview.short_description = "Photo"


@admin.register(AbstractThematicArea)
class AbstractThematicAreaAdmin(admin.ModelAdmin):
    list_display = ['name', 'conference', 'order', 'is_active']
    list_filter = ['conference', 'is_active']
    list_editable = ['order', 'is_active']


class AbstractSubmissionAdminForm(forms.ModelForm):
    reviewer_notes = forms.CharField(required=False, widget=CKEditorWidget(config_name='default'))
    class Meta:
        model = AbstractSubmission
        fields = '__all__'


@admin.register(AbstractSubmission)
class AbstractSubmissionAdmin(admin.ModelAdmin):
    form = AbstractSubmissionAdminForm
    list_display = ['reference_number', 'title_short', 'author_name', 'institution', 'thematic_area', 'presentation_format', 'status', 'word_count_display', 'submitted_at']
    list_filter = ['conference', 'status', 'thematic_area', 'presentation_format']
    search_fields = ['reference_number', 'title', 'author_name', 'institution', 'email']
    list_editable = ['status']
    readonly_fields = ['reference_number', 'submitted_at', 'updated_at', 'word_count_display']
    ordering = ['-submitted_at']

    fieldsets = (
        ('Reference', {
            'fields': ('reference_number', 'conference', 'submitted_at')
        }),
        ('Author', {
            'fields': ('author_name', 'co_authors', 'institution', 'email', 'phone')
        }),
        ('Abstract', {
            'fields': ('title', 'thematic_area', 'abstract_text', 'keywords', 'presentation_format', 'abstract_file')
        }),
        ('Review', {
            'fields': ('status', 'score', 'reviewer_notes')
        }),
    )

    def title_short(self, obj):
        return obj.title[:60] + ('…' if len(obj.title) > 60 else '')
    title_short.short_description = "Title"

    def word_count_display(self, obj):
        wc = obj.word_count()
        color = 'green' if wc <= 300 else 'red'
        return format_html('<span style="color:{}">{} words</span>', color, wc)
    word_count_display.short_description = "Word Count"

    actions = ['mark_under_review', 'mark_accepted', 'mark_rejected']

    def mark_under_review(self, request, queryset):
        queryset.update(status='under_review')
    mark_under_review.short_description = "Mark selected as Under Review"

    def mark_accepted(self, request, queryset):
        queryset.update(status='accepted')
    mark_accepted.short_description = "Mark selected as Accepted"

    def mark_rejected(self, request, queryset):
        queryset.update(status='rejected')
    mark_rejected.short_description = "Mark selected as Rejected"


@admin.register(RegistrationCategory)
class RegistrationCategoryAdmin(admin.ModelAdmin):
    list_display = ['name', 'conference', 'fee', 'early_bird_fee', 'max_slots', 'order', 'is_active']
    list_filter = ['conference', 'is_active']
    list_editable = ['fee', 'early_bird_fee', 'order', 'is_active']


@admin.register(Registration)
class RegistrationAdmin(ExportMixin, admin.ModelAdmin):
    resource_class = RegistrationResource
    list_display = ['ticket_id', 'full_name', 'email', 'category', 'amount', 'payment_method', 'payment_status',
                    'attendance_mode', 'attendance_response', 'is_stakeholder', 'checked_in', 'registered_at']
    list_filter = ['conference', 'is_stakeholder', 'payment_method', 'payment_status', 'attendance_mode',
                   'attendance_response', 'category', 'checked_in']
    search_fields = ['ticket_id', 'first_name', 'last_name', 'email', 'organization']
    list_editable = ['payment_status', 'checked_in']
    readonly_fields = ['ticket_id', 'registered_at', 'updated_at', 'paystack_reference', 'paystack_transaction_id',
                       'attendance_responded_at', 'attendance_email_sent_at', 'virtual_link_sent_at']
    ordering = ['-registered_at']

    fieldsets = (
        ('Ticket', {
            'fields': ('ticket_id', 'conference', 'category', 'registered_at')
        }),
        ('Participant', {
            'fields': ('first_name', 'last_name', 'email', 'phone', 'organization', 'position', 'state_of_origin')
        }),
        ('Preferences', {
            'fields': ('dietary_requirements', 't_shirt_size', 'abstract_reference'),
            'classes': ('collapse',),
        }),
        ('Payment', {
            'fields': ('is_stakeholder', 'amount', 'payment_method', 'payment_status', 'paystack_reference', 'paystack_transaction_id', 'payment_date')
        }),
        ('Attendance', {
            'fields': ('attendance_mode', 'attendance_response', 'attendance_responded_at', 'attendance_email_sent_at',
                       'virtual_link_sent_at')
        }),
        ('Accreditation', {
            'fields': ('checked_in', 'checked_in_at')
        }),
    )

    actions = ['export_selected_to_csv', 'confirm_registration',
               'resend_confirmation_emails', 'resend_payment_receipt',
               'send_attendance_confirmation', 'send_virtual_link', 'mark_checked_in']

    @admin.action(description='Send virtual joining link to selected')
    def send_virtual_link(self, request, queryset):
        from django.core.mail import get_connection
        from .emails import send_virtual_joining_link
        eligible = queryset.filter(payment_status__in=Registration.TICKET_VALID_STATUSES, attendance_mode='virtual')
        sent = failed = 0
        with get_connection() as connection:
            for registration in eligible.select_related('conference'):
                if not registration.conference.virtual_join_url:
                    failed += 1
                    continue
                try:
                    send_virtual_joining_link(registration, connection)
                except Exception:
                    failed += 1
                    continue
                Registration.objects.filter(pk=registration.pk).update(virtual_link_sent_at=timezone.now())
                sent += 1
        skipped = queryset.count() - eligible.count()
        msg = f"Joining link sent to {sent} virtual attendee(s)."
        if skipped:
            msg += f" {skipped} skipped (not virtual, or payment not confirmed)."
        self.message_user(request, msg, messages.SUCCESS if sent else messages.WARNING)
        if failed:
            self.message_user(request, f"{failed} not sent — check the joining link is set on the conference "
                                       "and the email logs.", messages.ERROR)

    @admin.action(description='Send physical attendance confirmation email to selected')
    def send_attendance_confirmation(self, request, queryset):
        from django.core.mail import get_connection
        from .emails import send_attendance_confirmation_request
        eligible = queryset.filter(payment_status__in=Registration.TICKET_VALID_STATUSES)
        sent = failed = 0
        with get_connection() as connection:
            for registration in eligible.select_related('conference'):
                url = request.build_absolute_uri(registration.attendance_confirmation_path())
                try:
                    send_attendance_confirmation_request(registration, url, connection=connection)
                except Exception:
                    failed += 1
                    continue
                Registration.objects.filter(pk=registration.pk).update(attendance_email_sent_at=timezone.now())
                sent += 1
        skipped = queryset.count() - eligible.count()
        msg = f"Attendance confirmation email sent to {sent} participant(s)."
        if skipped:
            msg += f" {skipped} skipped (payment not confirmed)."
        self.message_user(request, msg, messages.SUCCESS if sent else messages.WARNING)
        if failed:
            self.message_user(request, f"{failed} email(s) failed — check the email logs.", messages.ERROR)

    @admin.action(description='⬇ Export selected registrations to CSV')
    def export_selected_to_csv(self, request, queryset):
        """Download the selected registrations as a CSV file.

        The 'Export' button above exports the whole (filtered) list and offers
        other formats; this action is the quick path for a hand-picked set.
        """
        import csv
        from django.http import HttpResponse

        resource = self.resource_class()
        dataset = resource.export(queryset)

        response = HttpResponse(content_type='text/csv')
        filename = f"conference-registrations-{timezone.now():%Y%m%d-%H%M}.csv"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'

        writer = csv.writer(response)
        writer.writerow(dataset.headers)
        for row in dataset:
            writer.writerow(row)
        return response

    def confirm_registration(self, request, queryset):
        from django.utils import timezone
        # Save each instance individually (instead of queryset.update) so the
        # post_save signal fires and the receipt + welcome emails are sent.
        confirmed = 0
        for registration in queryset.exclude(payment_status='confirmed'):
            registration.payment_status = 'confirmed'
            if not registration.payment_date:
                registration.payment_date = timezone.now()
            registration.save()
            confirmed += 1
        self.message_user(
            request,
            f"{confirmed} registration(s) confirmed. Receipt and welcome emails sent.",
        )
    confirm_registration.short_description = "Confirm selected registrations"

    def resend_confirmation_emails(self, request, queryset):
        """(Re)send the receipt + welcome email to selected confirmed registrations."""
        from conference.emails import send_registration_confirmation
        confirmed = queryset.filter(payment_status='confirmed')
        sent = 0
        failed = 0
        for registration in confirmed:
            try:
                send_registration_confirmation(registration)
            except Exception:
                failed += 1
                continue
            Registration.objects.filter(pk=registration.pk).update(
                confirmation_email_sent=True, receipt_email_sent=True,
            )
            sent += 1

        skipped = queryset.count() - confirmed.count()
        msg = f"Receipt + welcome email sent to {sent} registration(s)."
        if skipped:
            msg += f" {skipped} skipped (payment not confirmed)."
        if failed:
            msg += f" {failed} failed — check the logs."
        self.message_user(request, msg, level=messages.WARNING if failed else messages.INFO)
    resend_confirmation_emails.short_description = "Resend receipt + welcome email"

    def resend_payment_receipt(self, request, queryset):
        """Send only the payment receipt to selected confirmed registrations."""
        from conference.emails import send_payment_receipt
        confirmed = queryset.filter(payment_status='confirmed')
        sent = 0
        failed = 0
        for registration in confirmed:
            try:
                send_payment_receipt(registration)
            except Exception:
                failed += 1
                continue
            Registration.objects.filter(pk=registration.pk).update(receipt_email_sent=True)
            sent += 1

        skipped = queryset.count() - confirmed.count()
        msg = f"Payment receipt sent to {sent} registration(s)."
        if skipped:
            msg += f" {skipped} skipped (payment not confirmed)."
        if failed:
            msg += f" {failed} failed — check the logs."
        self.message_user(request, msg, level=messages.WARNING if failed else messages.INFO)
    resend_payment_receipt.short_description = "Resend payment receipt only"

    def mark_checked_in(self, request, queryset):
        from django.utils import timezone
        queryset.update(checked_in=True, checked_in_at=timezone.now())
    mark_checked_in.short_description = "Mark selected as Checked In"

    def full_name(self, obj):
        return obj.full_name
    full_name.short_description = "Name"


class ProgramSessionAdminForm(forms.ModelForm):
    description = forms.CharField(required=False, widget=CKEditorWidget(config_name='default'))
    class Meta:
        model = ProgramSession
        fields = '__all__'


class ProgramSessionInline(admin.StackedInline):
    model = ProgramSession
    form = ProgramSessionAdminForm
    extra = 1
    fields = ['title', 'session_type', 'start_time', 'end_time', 'venue', 'speakers', 'moderator', 'description', 'order', 'is_active']
    filter_horizontal = ['speakers']


@admin.register(ProgramDay)
class ProgramDayAdmin(admin.ModelAdmin):
    list_display = ['title', 'conference', 'date', 'theme', 'order', 'is_active']
    list_filter = ['conference', 'is_active']
    list_editable = ['order', 'is_active']
    inlines = [ProgramSessionInline]


@admin.register(ProgramSession)
class ProgramSessionAdmin(admin.ModelAdmin):
    form = ProgramSessionAdminForm
    list_display = ['title', 'day', 'session_type', 'start_time', 'end_time', 'venue', 'order']
    list_filter = ['day__conference', 'session_type']
    search_fields = ['title', 'moderator']
    filter_horizontal = ['speakers']
    ordering = ['day', 'order', 'start_time']


@admin.register(Sponsor)
class SponsorAdmin(admin.ModelAdmin):
    list_display = ['name', 'conference', 'tier', 'order', 'is_active']
    list_filter = ['conference', 'tier', 'is_active']
    list_editable = ['tier', 'order', 'is_active']


class KeyMessageAdminForm(forms.ModelForm):
    message = forms.CharField(widget=CKEditorWidget(config_name='default'))
    class Meta:
        model = KeyMessage
        fields = '__all__'


@admin.register(KeyMessage)
class KeyMessageAdmin(admin.ModelAdmin):
    form = KeyMessageAdminForm
    list_display = ['message_short', 'conference', 'source', 'is_quote', 'order', 'is_active']
    list_filter = ['conference', 'is_quote', 'is_active']
    list_editable = ['order', 'is_active']

    def message_short(self, obj):
        from django.utils.html import strip_tags
        return strip_tags(obj.message)[:80]
    message_short.short_description = "Message"


@admin.register(LOCMember)
class LOCMemberAdmin(admin.ModelAdmin):
    list_display = ['full_name', 'position', 'email', 'phone', 'conference', 'order', 'is_active']
    list_filter = ['conference', 'is_active']
    search_fields = ['full_name', 'position', 'email']
    list_editable = ['order', 'is_active']
    ordering = ['conference', 'order']


@admin.register(ExhibitorPackage)
class ExhibitorPackageAdmin(admin.ModelAdmin):
    """Standalone exhibition booths that can be booked and paid for."""
    list_display = ['name', 'conference', 'price', 'space_structure', 'max_slots', 'order', 'is_active']
    list_filter = ['conference', 'is_active']
    search_fields = ['name', 'space_structure', 'perks']
    list_editable = ['price', 'order', 'is_active']
    ordering = ['conference', 'order']

    fieldsets = (
        ('Booth', {
            'fields': ('conference', 'name', 'description', 'icon')
        }),
        ('Price & space', {
            'fields': ('price', 'space_structure'),
            'description': 'Space/structure appears beside the price on the Exhibition page.'
        }),
        ('What is included', {
            'fields': ('perks',),
            'description': 'One item per line. Each line becomes a ticked bullet on the booth card.'
        }),
        ('Availability', {
            'fields': ('max_slots', 'order', 'is_active')
        }),
    )


class ExhibitorShowcaseInline(admin.TabularInline):
    model = ExhibitorShowcase
    extra = 0
    fields = ['image_preview', 'title', 'price', 'is_approved', 'is_active', 'created_at']
    readonly_fields = ['image_preview', 'created_at']

    def image_preview(self, obj):
        if obj.image:
            return format_html('<img src="{}" style="width:60px;height:60px;object-fit:cover;border-radius:8px;">', obj.image.url)
        return "—"
    image_preview.short_description = "Image"


@admin.register(Exhibitor)
class ExhibitorAdmin(admin.ModelAdmin):
    list_display = ['reference', 'company_name', 'contact_name', 'email', 'package', 'amount', 'payment_method', 'payment_status', 'registered_at']
    list_filter = ['conference', 'payment_method', 'payment_status', 'package']
    search_fields = ['reference', 'company_name', 'contact_name', 'email']
    list_editable = ['payment_status']
    readonly_fields = ['reference', 'access_token', 'registered_at', 'updated_at', 'paystack_reference', 'paystack_transaction_id', 'showcase_link']
    ordering = ['-registered_at']
    inlines = [ExhibitorShowcaseInline]

    fieldsets = (
        ('Reference', {
            'fields': ('reference', 'conference', 'package', 'access_token', 'showcase_link', 'registered_at')
        }),
        ('Exhibitor', {
            'fields': ('company_name', 'contact_name', 'email', 'phone', 'website', 'logo')
        }),
        ('Payment', {
            'fields': ('amount', 'payment_method', 'payment_status', 'paystack_reference', 'paystack_transaction_id', 'payment_date')
        }),
    )

    actions = ['confirm_exhibitor']

    def confirm_exhibitor(self, request, queryset):
        queryset.update(payment_status='confirmed', payment_date=timezone.now())
    confirm_exhibitor.short_description = "Confirm selected exhibitors (payment)"

    def showcase_link(self, obj):
        if obj.pk:
            from django.urls import reverse
            url = reverse('conference:exhibitor_showcase', args=[obj.access_token])
            return format_html('<a href="{}" target="_blank">{}</a>', url, url)
        return "—"
    showcase_link.short_description = "Private showcase link"


@admin.register(ExhibitorShowcase)
class ExhibitorShowcaseAdmin(admin.ModelAdmin):
    list_display = ['title', 'exhibitor', 'price', 'is_approved', 'is_active', 'image_preview', 'created_at']
    list_filter = ['is_approved', 'is_active', 'exhibitor__conference']
    search_fields = ['title', 'exhibitor__company_name']
    list_editable = ['is_approved', 'is_active']
    readonly_fields = ['created_at', 'image_preview']
    ordering = ['-created_at']

    actions = ['approve_items', 'unapprove_items']

    def approve_items(self, request, queryset):
        queryset.update(is_approved=True)
    approve_items.short_description = "Approve selected items for public display"

    def unapprove_items(self, request, queryset):
        queryset.update(is_approved=False)
    unapprove_items.short_description = "Unapprove selected items"

    def image_preview(self, obj):
        if obj.image:
            return format_html('<img src="{}" style="width:60px;height:60px;object-fit:cover;border-radius:8px;">', obj.image.url)
        return "—"
    image_preview.short_description = "Image"


@admin.register(PaymentVerifier)
class PaymentVerifierAdmin(admin.ModelAdmin):
    list_display = ['name', 'email', 'is_active', 'last_login_at', 'created_at']
    list_filter = ['is_active']
    search_fields = ['name', 'email']
    list_editable = ['is_active']
    readonly_fields = ['last_login_at', 'created_at']
    fieldsets = (
        (None, {
            'fields': ('name', 'email', 'is_active'),
            'description': (
                'People listed here can access the Payment Verification page without being '
                'staff. They visit /conference/payment/verification/, enter this email, and '
                'receive a 2-hour sign-in link. Untick "is active" to revoke access immediately.'
            ),
        }),
        ('Activity', {
            'fields': ('last_login_at', 'created_at'),
        }),
    )


@admin.register(AccreditationOfficer)
class AccreditationOfficerAdmin(admin.ModelAdmin):
    list_display = ['name', 'email', 'is_active', 'last_login_at', 'created_at']
    list_filter = ['is_active']
    search_fields = ['name', 'email']
    list_editable = ['is_active']
    readonly_fields = ['last_login_at', 'created_at']
    fieldsets = (
        (None, {
            'fields': ('name', 'email', 'is_active'),
            'description': (
                'People listed here can use the Accreditation page at the venue entrance, '
                'without access to payment verification. They visit /conference/accreditation/, '
                'enter this email, and receive a 2-hour sign-in link. Untick "is active" to '
                'revoke access immediately.'
            ),
        }),
        ('Activity', {
            'fields': ('last_login_at', 'created_at'),
        }),
    )


@admin.register(AbstractReviewer)
class AbstractReviewerAdmin(admin.ModelAdmin):
    list_display = ['name', 'email', 'is_active', 'link_sent_at', 'last_login_at', 'created_at']
    list_filter = ['is_active']
    search_fields = ['name', 'email']
    list_editable = ['is_active']
    readonly_fields = ['access_link', 'link_sent_at', 'last_login_at', 'created_at']
    actions = ['resend_access_link']
    fieldsets = (
        (None, {
            'fields': ('name', 'email', 'is_active'),
            'description': (
                'People listed here can view the Abstract Review page without being staff. '
                'When you add an active email and save, their permanent access link is emailed '
                'to them automatically. The link never expires — they can bookmark it and return '
                'any time while active. Copy it from "Access link" below to share it another way, '
                'use the "Resend access link" action to email it again, or untick "is active" to '
                'revoke access immediately.'
            ),
        }),
        ('Access', {
            'fields': ('access_link', 'link_sent_at', 'last_login_at', 'created_at'),
        }),
    )

    def changeform_view(self, request, *args, **kwargs):
        # Stash the request so access_link can build an absolute URL.
        self._request = request
        return super().changeform_view(request, *args, **kwargs)

    def access_link(self, obj):
        if not obj or not obj.pk:
            return "— saved after you add the reviewer —"
        from .views import abstract_reviewer_access_url
        request = getattr(self, '_request', None)
        if request is not None:
            url = abstract_reviewer_access_url(request, obj)
        else:
            from django.urls import reverse
            url = reverse('conference:abstract_reviewer_login', args=[str(obj.access_token)])
        return format_html(
            '<a href="{}" target="_blank" style="word-break:break-all;">{}</a>'
            '<br><span style="color:#888;font-size:0.85em;">Permanent — share only with this reviewer.</span>',
            url, url,
        )
    access_link.short_description = "Access link"

    def _send_link(self, request, reviewer):
        """Send a magic link to one reviewer; report success/failure via messages.
        Returns True on success."""
        from .views import issue_abstract_reviewer_link
        if not reviewer.is_active:
            messages.warning(
                request,
                f"Skipped {reviewer.email}: access is deactivated (tick “is active” first).",
            )
            return False
        try:
            issue_abstract_reviewer_link(request, reviewer)
        except Exception as exc:
            messages.error(request, f"Could not email {reviewer.email}: {exc}")
            return False
        messages.success(request, f"Access link sent to {reviewer.email}.")
        return True

    def save_model(self, request, obj, form, change):
        super().save_model(request, obj, form, change)
        # Email an access link when a reviewer is first added (or when an
        # existing entry is (re)activated and has never received one).
        newly_active = 'is_active' in form.changed_data and obj.is_active
        if obj.is_active and (not change or newly_active) and not obj.link_sent_at:
            self._send_link(request, obj)

    @admin.action(description="Resend access link")
    def resend_access_link(self, request, queryset):
        for reviewer in queryset:
            self._send_link(request, reviewer)


@admin.register(ContentBlock)
class ContentBlockAdmin(admin.ModelAdmin):
    list_display = ['key', 'content_preview', 'updated_by', 'updated_at']
    search_fields = ['key', 'content']
    readonly_fields = ['updated_by', 'updated_at']

    def content_preview(self, obj):
        from django.utils.html import strip_tags
        return strip_tags(obj.content)[:80]
    content_preview.short_description = "Content"


# ─── Sponsorship packages & benefits ──────────────────────────────────────────

@admin.register(SponsorshipBenefit)
class SponsorshipBenefitAdmin(admin.ModelAdmin):
    """The rows of the sponsorship comparison table."""
    list_display = ['label', 'conference', 'order', 'packages_filled', 'is_active']
    list_editable = ['order', 'is_active']
    list_filter = ['conference', 'is_active']
    search_fields = ['label', 'description']
    ordering = ['conference', 'order']

    fieldsets = (
        (None, {
            'fields': ('conference', 'label', 'description', 'order', 'is_active'),
            'description': (
                'Each benefit is one <strong>row</strong> of the comparison table on the '
                'Exhibition page. What each package offers for this row is set on the '
                'package itself, under Sponsorship Packages.'
            )
        }),
    )

    @admin.display(description='Filled in for')
    def packages_filled(self, obj):
        total = SponsorshipPackage.objects.filter(conference=obj.conference, is_active=True).count()
        filled = obj.package_values.exclude(value='').count()
        colour = '#0a7d2b' if filled >= total else '#b45309'
        return format_html('<span style="color:{};">{} of {} package(s)</span>', colour, filled, total)


class SponsorshipPackageBenefitInline(admin.TabularInline):
    """Edit every benefit value for a package on one screen."""
    model = SponsorshipPackageBenefit
    extra = 0
    fields = ['benefit', 'value']
    autocomplete_fields = ['benefit']
    verbose_name = "Benefit"
    verbose_name_plural = "What this package includes (one row per benefit)"

    def get_formset(self, request, obj=None, **kwargs):
        formset = super().get_formset(request, obj, **kwargs)
        if obj:
            # Only offer benefits belonging to this package's conference.
            formset.form.base_fields['benefit'].queryset = SponsorshipBenefit.objects.filter(
                conference=obj.conference, is_active=True
            )
        return formset


@admin.register(SponsorshipPackage)
class SponsorshipPackageAdmin(admin.ModelAdmin):
    """The columns of the sponsorship comparison table."""
    list_display = ['name', 'conference', 'price_column', 'position', 'benefits_filled',
                    'is_featured', 'order', 'is_active']
    list_editable = ['is_featured', 'order', 'is_active']
    list_filter = ['conference', 'is_active', 'is_featured']
    search_fields = ['name', 'position']
    ordering = ['conference', 'order']
    inlines = [SponsorshipPackageBenefitInline]

    fieldsets = (
        ('Package', {
            'fields': ('conference', 'name', 'position'),
        }),
        ('Price', {
            'fields': ('price', 'price_display'),
            'description': (
                'Enter the amount in <strong>Price</strong> for a fixed tier. For a range '
                '(such as the Supporting Partner tier) leave Price blank and type the range '
                'into <strong>Price display</strong> instead.'
            )
        }),
        ('Display', {
            'fields': ('accent_color', 'is_featured', 'order', 'is_active'),
        }),
    )

    @admin.display(description='Price', ordering='price')
    def price_column(self, obj):
        return obj.display_price

    @admin.display(description='Benefits filled')
    def benefits_filled(self, obj):
        total = SponsorshipBenefit.objects.filter(conference=obj.conference, is_active=True).count()
        filled = obj.benefit_values.exclude(value='').count()
        if not total:
            return format_html('<span style="color:#b45309;">no benefit rows defined</span>')
        colour = '#0a7d2b' if filled >= total else '#b45309'
        return format_html('<span style="color:{};">{} of {}</span>', colour, filled, total)


@admin.register(SponsorshipPackageBenefit)
class SponsorshipPackageBenefitAdmin(admin.ModelAdmin):
    """Direct access to individual cells, for bulk edits across packages."""
    list_display = ['benefit', 'package', 'short_value']
    list_filter = ['package__conference', 'benefit', 'package']
    search_fields = ['value', 'benefit__label', 'package__name']
    autocomplete_fields = ['package', 'benefit']
    ordering = ['benefit__order', 'package__order']

    @admin.display(description='Value')
    def short_value(self, obj):
        if not obj.is_included:
            return format_html('<span style="color:#999;">not included</span>')
        return obj.value[:80] + ('…' if len(obj.value) > 80 else '')
